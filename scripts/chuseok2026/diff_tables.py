"""새로 받은 추석 행로표 엑셀(보안 없는 것) → 지난번에 넣은 자료와 다이아별로 대조.

쓰는 법(저장소 뿌리에서): py -3 scripts/chuseok2026/diff_tables.py <엑셀 파일> <비교 대상>
  비교 대상: ph(9/23 평휴 연결) · hol(주간) · hh(휴휴 9/24~25) · hh0926(9/26 심야연장) · hp0927(9/27 심야연장)

비교하는 것:
  · 다이아 번호 목록(생기거나 없어진 번호)
  · 근무시간 계 · 운전 · 편승 · 주행키로 (엑셀 요약칸)
  · 행로표 그림 — 픽셀이 같은지(같으면 내용이 같다)
그림이 다르거나 요약값이 다른 다이아만 다시 옮겨 적으면 된다.

지난 자료는 저장소 안에서 찾는다 — 요약값은 이 폴더의 xl_summary.json,
그림은 앱이 쓰는 public/images/route/chuseok2026/ (엑셀 그림 그대로다).
9/23 연결시작 그림은 캡처(jpg)라 엑셀 그림과 모양이 달라 그림 비교는 건너뛰고 요약값만 본다.

⚠ 행로표 시스템에서 받은 «보안 없는» 엑셀이어야 한다. 사업소가 돌린 문서보안(MarkAny) 파일은
   파일 머리가 <DOCUMENTSAFER 로 시작하고, 이 PC 에서는 엑셀로도 열리지 않는다(2026-09-21 확인).
"""
import openpyxl, sys, os, io, json, hashlib, re
from PIL import Image, ImageChops

LABELS = ['계', '운전', '준비', '대기', '감시', '입환', '편승', '정리', '야간']
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
IMG_ROOT = os.path.join(ROOT, 'public', 'images', 'route', 'chuseok2026')
OLD_SUMMARY = json.load(io.open(os.path.join(HERE, 'xl_summary.json'), encoding='utf-8'))
# 지난번 표 → (요약 키, 그림 파일 앞글자). 그림 앞글자가 None 이면 그림은 비교하지 않는다.
OLD_MAP = {
    'ph': ('ph', None),          # 9/23 연결시작 — 앱 그림은 캡처라 엑셀 그림과 비교할 수 없다
    'hol': ('hol', 'day'),
    'hh': ('hh', 'hh'),
    'hh0926': ('hh', 'hh'),      # 지난번엔 9/26 = 휴휴 + 심야연장 변경분(캡처)로 만들었다
    'hp0927': ('hp', 'hp'),      # 지난번엔 9/27 = 휴평 + 심야연장 변경분(캡처)로 만들었다
}
#: 표 머리의 «NO» 같은 칸은 다이아가 아니다
def is_dia(k: str) -> bool:
    return bool(re.fullmatch(r'(?:임시|변)?\d+', k))


def dia_name(v) -> str:
    """칸의 다이아 번호 — «임시\\n61» 처럼 줄이 바뀐 것도 «임시61» 로"""
    return str(int(v)) if isinstance(v, (int, float)) else re.sub(r'\s+', '', str(v))


def read_workbook(path):
    wb = openpyxl.load_workbook(path)
    wbv = openpyxl.load_workbook(path, data_only=True)
    summ, imgs = {}, {}
    for ws, wsv in zip(wb.worksheets, wbv.worksheets):
        g = {(c.row, c.column): c.value for row in wsv.iter_rows() for c in row if c.value not in (None, '')}
        labels = []   # (행, 열, 다이아) — 다이아 번호 칸의 자리
        for (r, c), v in g.items():
            if isinstance(v, str) and v.replace(' ', '') == '승무구간':
                dv = g.get((r, c - 2))
                if dv is None:
                    continue
                dia = dia_name(dv)
                info = {}
                for (r2, c2), v2 in g.items():
                    if r < r2 <= r + 30 and c2 == c + 30 and v2 in LABELS:
                        info[v2] = g.get((r2, c2 + 2))
                    if r < r2 <= r + 30 and c2 == c and v2 == '주행키로':
                        info['km'] = g.get((r2, c2 + 4))
                if is_dia(dia):
                    summ[dia] = info
                    labels.append((r, c - 2, dia))
        for im in ws._images:
            a = im.anchor._from
            row, col = a.row + 1, a.col + 1
            # 그림은 보통 번호 칸에서 2행 아래·2열 오른쪽에 붙는데, 손으로 갈아 끼운 그림은 자리가 조금씩
            # 어긋난다(9/27 심야연장 반영본에서 row 2·col 40 등). 그림의 왼쪽 위에서 가장 가까운 번호 칸을 쓴다.
            near = [(r, c, d) for r, c, d in labels if r <= row + 2 and c <= col + 2]
            if near:
                imgs[max(near)[2]] = im._data()
    return summ, imgs


def same_picture(new_bytes, old_path):
    if not os.path.exists(old_path):
        return False
    a = Image.open(io.BytesIO(new_bytes)).convert('RGB')
    b = Image.open(old_path).convert('RGB')
    if a.size != b.size:
        return False
    return ImageChops.difference(a, b).getbbox() is None


def main():
    path, kind = sys.argv[1], sys.argv[2]
    key, img_prefix = OLD_MAP[kind]
    old = {k: v for k, v in OLD_SUMMARY[key].items() if is_dia(k)}
    new, imgs = read_workbook(path)
    out_dir = os.path.join(HERE, 'new', kind)  # 새 그림 — 바뀐 다이아를 옮겨 적을 때 본다
    os.makedirs(out_dir, exist_ok=True)
    for dia, b in imgs.items():
        open(os.path.join(out_dir, f'{dia}.png'), 'wb').write(b)

    num = lambda k: (0, int(k)) if k.isdigit() else (1, k)
    added = sorted(set(new) - set(old), key=num)
    removed = sorted(set(old) - set(new), key=num)
    changed = []
    for dia in sorted(set(new) & set(old), key=num):
        diffs = [f'{f} {old[dia].get(f)}→{new[dia].get(f)}' for f in ('계', '운전', '편승', 'km')
                 if str(old[dia].get(f)) != str(new[dia].get(f))]
        pic_same = img_prefix is None or (
            dia in imgs and same_picture(imgs[dia], os.path.join(IMG_ROOT, f'{img_prefix}_{dia}.png')))
        if diffs or not pic_same:
            changed.append((dia, diffs, pic_same))
    print(f'[{kind}] 새 파일 다이아 {len(new)}개 · 지난 자료 {len(old)}개')
    print('  새로 생김:', ' '.join(added) or '없음')
    print('  없어짐  :', ' '.join(removed) or '없음')
    print(f'  바뀜    : {len(changed)}개')
    for dia, diffs, pic_same in changed:
        print(f'    {dia}번 — ' + ('; '.join(diffs) if diffs else '요약값 같음') + ('' if pic_same else ' · 그림 다름'))
    json.dump({'added': added, 'removed': removed, 'changed': [c[0] for c in changed]},
              io.open(os.path.join(out_dir, 'diff.json'), 'w', encoding='utf-8'), ensure_ascii=False)


if __name__ == '__main__':
    main()
